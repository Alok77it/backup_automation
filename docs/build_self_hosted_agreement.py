from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("Self-Hosted_Software_License_Agreement.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E2EC")


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 15, "1F4D78", 14, 7),
        ("Heading 2", 12.5, "2E74B5", 10, 5),
        ("Heading 3", 11.5, "1F4D78", 8, 4),
    ]:
        s = styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("SELF-HOSTED SOFTWARE LICENSE AND DEPLOYMENT AGREEMENT")
    run.bold = True
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(11, 37, 69)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("For proprietary deployment of Backup Automation / Backup Intelligence software")
    run.italic = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(75, 85, 99)

    table = doc.add_table(rows=4, cols=2)
    table.autofit = False
    set_table_borders(table)
    widths = [Inches(2.0), Inches(4.7)]
    labels = [
        ("Agreement Date", "[Effective Date]"),
        ("Licensor / Provider", "Alok Trivedi / [Provider Legal Entity]"),
        ("Customer / Licensee", "[Customer Legal Name]"),
        ("Authorized Deployment", "[Server / Domain / Organization / Environment]"),
    ]
    for row_idx, (label, value) in enumerate(labels):
        for col_idx, width in enumerate(widths):
            table.cell(row_idx, col_idx).width = width
            set_cell_margins(table.cell(row_idx, col_idx))
            table.cell(row_idx, col_idx).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(table.cell(row_idx, 0), "F2F4F7")
        table.cell(row_idx, 0).paragraphs[0].add_run(label).bold = True
        table.cell(row_idx, 1).paragraphs[0].add_run(value)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(10)
    note.paragraph_format.left_indent = Inches(0.15)
    note.paragraph_format.right_indent = Inches(0.15)
    r = note.add_run("Important: ")
    r.bold = True
    note.add_run(
        "This template is a business agreement draft. It should be reviewed by legal counsel before signature or customer use."
    )


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def make_doc():
    doc = Document()
    style_document(doc)

    header = doc.sections[0].header.paragraphs[0]
    header.text = "Self-Hosted Software License and Deployment Agreement"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(107, 114, 128)

    add_title(doc)

    h(doc, "1. Parties and Purpose")
    para(doc, "This Self-Hosted Software License and Deployment Agreement (the \"Agreement\") is entered into by and between Alok Trivedi / [Provider Legal Entity] (the \"Provider\" or \"Licensor\") and [Customer Legal Name] (the \"Customer\" or \"Licensee\").")
    para(doc, "The purpose of this Agreement is to define the terms under which the Provider may install, deploy, operate, support, or license the proprietary Backup Automation / Backup Intelligence software and related services on infrastructure controlled by the Customer.")

    h(doc, "2. Definitions")
    definitions = [
        ("Software", "the application, source code, object code, backend logic, APIs, workers, frontend, database schema, deployment files, documentation, updates, scripts, designs, and related materials supplied or installed by the Provider."),
        ("Self-Hosted Deployment", "any installation or operation of the Software on servers, cloud accounts, networks, infrastructure, or environments owned, leased, controlled, or administered by the Customer or a third party acting for the Customer."),
        ("Authorized Environment", "the specific server, domain, customer organization, users, storage scope, and environment approved in writing by the Provider."),
        ("Confidential Information", "all non-public technical, commercial, operational, pricing, customer, security, source code, and deployment information related to the Software."),
    ]
    for term, meaning in definitions:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{term}: ").bold = True
        p.add_run(meaning)

    h(doc, "3. License Grant")
    para(doc, "Subject to payment and compliance with this Agreement, the Provider grants the Customer a limited, non-exclusive, non-transferable, revocable license to use the Software only within the Authorized Environment and only for the Customer's internal business operations.")
    para(doc, "No ownership or intellectual property rights are transferred. All rights not expressly granted remain reserved by the Provider.")

    h(doc, "4. Self-Hosted Deployment Conditions")
    bullet(doc, "Self-hosted use is available only under an enterprise or written commercial license.")
    bullet(doc, "Installation, configuration, updates, migrations, and maintenance must be performed by the Provider or with the Provider's written approval.")
    bullet(doc, "The Customer must not move, clone, duplicate, export, snapshot, or redeploy the Software to another server, domain, cloud account, or organization without written permission.")
    bullet(doc, "The Customer must not provide application file, repository, container, database, or administrative access to unauthorized persons.")
    bullet(doc, "The Customer is responsible for infrastructure availability, cloud costs, network access, physical security, and any third-party hosting charges unless otherwise stated in writing.")

    h(doc, "5. Source Code, Customization, and Modification Restrictions")
    para(doc, "The Customer may not access, copy, modify, customize, patch, translate, adapt, fork, or create derivative works from the Software unless the Provider gives prior written permission.")
    para(doc, "The Customer may not modify backend logic, API routes, worker logic, automation workflows, database schema, license controls, billing controls, deployment scripts, security controls, branding, or copyright notices.")
    para(doc, "Any customization requested by the Customer must be scoped, quoted, approved, and delivered by the Provider under a separate written statement of work or change order.")

    h(doc, "6. Prohibited Conduct")
    for item in [
        "reverse engineering, decompiling, disassembling, unpacking, extracting, or attempting to discover source code, architecture, algorithms, API logic, worker logic, or trade secrets;",
        "reselling, sublicensing, renting, leasing, redistributing, publishing, leaking, transferring, or sharing the Software or access to it;",
        "using the Software to build, operate, assist, or launch a competing product, service, platform, or managed offering;",
        "removing, hiding, changing, or bypassing license keys, activation checks, domain restrictions, usage limits, audit controls, telemetry controls, payment controls, or security controls;",
        "using the Software for white-label services, third-party hosting, marketplace listings, or managed services without written permission;",
        "attempting to avoid fees, usage limits, contractual restrictions, or commercial revenue-share obligations.",
    ]:
        bullet(doc, item)

    h(doc, "7. Commercial Use, Resale, and Revenue Share")
    para(doc, "Enterprise deployment, resale, redistribution, white-label use, third-party hosting, managed-service use, marketplace listing, bundling, or any use that generates direct or indirect revenue requires a separate written commercial agreement.")
    para(doc, "If the Customer or any related party uses the Software, a modified version, a copied version, a reverse-engineered version, or a substantially similar derivative for enterprise deployment, resale, redistribution, white-label service, managed service, or commercial sale without written authorization, the Customer agrees to pay the Provider forty percent (40%) of gross revenue, contract value, subscription fees, service fees, resale proceeds, or other commercial benefit derived from such use.")
    para(doc, "The Provider may agree in writing to a different percentage, including a thirty percent (30%) to forty percent (40%) range, but no reduced percentage applies unless signed by the Provider.")

    h(doc, "8. Fees, Payment, and Taxes")
    bullet(doc, "The Customer must pay all setup fees, license fees, support fees, customization fees, infrastructure charges, storage charges, usage charges, and applicable taxes stated in the order form, invoice, proposal, or statement of work.")
    bullet(doc, "Payment delays may result in suspension of support, updates, access, license validation, or continued use.")
    bullet(doc, "All fees are non-refundable unless expressly stated otherwise in writing.")

    h(doc, "9. Access, Administration, and Security")
    para(doc, "The Provider may require administrator or root access for installation, updates, maintenance, troubleshooting, backup verification, migrations, and security hardening. Such access must be used only for service delivery and must be handled with reasonable care.")
    para(doc, "The Customer must not provide direct dashboard, server, repository, database, container, or operating system access to unauthorized persons. The Customer must maintain strong credentials, restrict access by role, and promptly report suspected compromise.")
    para(doc, "If the Customer changes server files, permissions, containers, images, credentials, network rules, database state, or deployment configuration without approval, support may be suspended and remediation may be billable.")

    h(doc, "10. Support, Updates, and Maintenance")
    para(doc, "Support, updates, migrations, backups, monitoring, security patches, and maintenance are provided only if included in a signed agreement, proposal, invoice, or active support plan.")
    para(doc, "The Provider is not responsible for failures caused by unauthorized modification, unsupported infrastructure, third-party service outages, customer network policies, customer credentials, deleted files, expired domains, insufficient storage, or non-payment.")

    h(doc, "11. Data, Backups, and Infrastructure")
    bullet(doc, "The Customer remains responsible for the legality, accuracy, ownership, retention, and protection of its data.")
    bullet(doc, "Backup schedules, retention periods, recovery objectives, storage capacity, and monitoring requirements must be agreed in writing.")
    bullet(doc, "The Provider is not liable for data loss unless expressly stated in a signed written agreement.")

    h(doc, "12. Confidentiality")
    para(doc, "Each party must protect the other party's Confidential Information using reasonable safeguards. The Customer must not disclose, publish, leak, or permit unauthorized access to the Software, source code, deployment files, credentials, pricing, architecture, or non-public technical information.")
    para(doc, "Confidentiality obligations survive termination of this Agreement.")

    h(doc, "13. Audit and Verification")
    para(doc, "The Provider may request reasonable verification of server count, domain count, user count, customer count, storage use, usage volume, deployment scope, commercial use, and revenue related to the Software. Refusal to provide reasonable verification may be treated as a material breach.")

    h(doc, "14. Intellectual Property")
    para(doc, "The Software, all updates, improvements, bug fixes, customizations, derivatives, documentation, designs, methods, workflows, and know-how remain the exclusive property of the Provider unless a signed written agreement states otherwise.")
    para(doc, "Customer feedback, suggestions, feature requests, and improvement ideas may be used by the Provider without restriction or compensation, unless prohibited by a signed written agreement.")

    h(doc, "15. Third-Party Components")
    para(doc, "The Software may include or depend on third-party libraries, frameworks, operating systems, containers, APIs, databases, or services. Those components remain subject to their own licenses and terms. This Agreement applies to the Provider's proprietary Software and services.")

    h(doc, "16. Suspension and Termination")
    para(doc, "The Provider may suspend or terminate the license, support, updates, or access if the Customer breaches this Agreement, fails to pay, exceeds the authorized scope, modifies the deployment without approval, tampers with license controls, reverse engineers the Software, redistributes the Software, or uses it commercially without authorization.")
    para(doc, "Upon termination, the Customer must stop using the Software and delete or return all copies, deployments, containers, images, credentials, documentation, and related materials unless otherwise agreed in writing.")

    h(doc, "17. Warranty Disclaimer")
    para(doc, "The Software is provided \"as is\" unless a signed written agreement states otherwise. The Provider disclaims all warranties to the maximum extent permitted by law, including warranties of merchantability, fitness for a particular purpose, non-infringement, uninterrupted operation, error-free operation, security, availability, and data loss prevention.")

    h(doc, "18. Limitation of Liability")
    para(doc, "To the maximum extent permitted by law, the Provider is not liable for indirect, incidental, special, consequential, punitive, or business-loss damages, including lost profits, lost revenue, downtime, data loss, security incidents, business interruption, or replacement service costs.")
    para(doc, "Unless a signed written agreement states otherwise, the Provider's total liability is limited to the fees paid by the Customer for the Software or services during the three (3) months before the event giving rise to the claim.")

    h(doc, "19. Remedies")
    para(doc, "Unauthorized copying, modification, reverse engineering, resale, redistribution, sublicensing, commercial exploitation, or disclosure may cause irreparable harm. The Provider may seek injunctive relief, damages, unpaid fees, revenue share, audit costs, legal costs, and any other remedies available under law or contract.")

    h(doc, "20. Governing Law and Dispute Resolution")
    para(doc, "This Agreement is governed by the laws of [Jurisdiction]. The courts or arbitration forum located in [Venue] will have exclusive jurisdiction unless the parties agree otherwise in writing.")

    h(doc, "21. Entire Agreement and Order of Precedence")
    para(doc, "This Agreement, together with any signed order form, proposal, invoice, statement of work, or written addendum, forms the entire agreement for the self-hosted deployment. If there is a conflict, the signed written agreement or order form controls for that specific conflict.")

    h(doc, "22. Signature")
    para(doc, "By signing below, each party confirms that it has authority to enter into this Agreement and agrees to be bound by its terms.")

    sig = doc.add_table(rows=5, cols=2)
    sig.autofit = False
    set_table_borders(sig)
    headers = ["Provider / Licensor", "Customer / Licensee"]
    for i, header_text in enumerate(headers):
        sig.cell(0, i).text = header_text
        sig.cell(0, i).paragraphs[0].runs[0].bold = True
        set_cell_shading(sig.cell(0, i), "F2F4F7")
        sig.cell(0, i).width = Inches(3.25)
    rows = ["Name:", "Title:", "Signature:", "Date:"]
    for r, label in enumerate(rows, start=1):
        for c in range(2):
            set_cell_margins(sig.cell(r, c), top=120, bottom=120)
            sig.cell(r, c).paragraphs[0].add_run(label).bold = True
            sig.cell(r, c).paragraphs[0].add_run(" ______________________________")

    doc.add_page_break()
    h(doc, "Schedule A: Commercial and Deployment Details")
    para(doc, "Complete this schedule for each customer deployment.")
    schedule = doc.add_table(rows=1, cols=3)
    schedule.autofit = False
    set_table_borders(schedule)
    set_repeat_table_header(schedule.rows[0])
    headers = ["Item", "Agreed Value", "Notes"]
    widths = [Inches(2.0), Inches(2.25), Inches(2.25)]
    for i, text in enumerate(headers):
        cell = schedule.cell(0, i)
        cell.text = text
        cell.width = widths[i]
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell)
        cell.paragraphs[0].runs[0].bold = True
    items = [
        "Authorized organization",
        "Authorized domain / URL",
        "Authorized server / VPS",
        "Deployment type",
        "Storage allowance",
        "User / agent limits",
        "Setup fee",
        "Monthly / annual license fee",
        "Maintenance / support fee",
        "Customization scope",
        "Payment terms",
        "Support hours / SLA",
        "Backup retention",
        "Special restrictions",
    ]
    for item in items:
        row = schedule.add_row()
        row.cells[0].text = item
        row.cells[1].text = "[To be completed]"
        row.cells[2].text = "[Notes]"
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "Confidential and Proprietary"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(107, 114, 128)

    doc.save(OUT)


if __name__ == "__main__":
    make_doc()
    print(OUT)
